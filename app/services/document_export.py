"""Render academic Markdown into downloadable DOCX and PDF files."""
from __future__ import annotations

import re
from io import BytesIO
from xml.sax.saxutils import escape


def _plain(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return re.sub(r"(`{1,3}|\*\*|__|~~)", "", text).strip()


def _blocks(markdown: str):
    """Yield a small, deterministic subset of Markdown as semantic blocks."""
    paragraph: list[str] = []

    def flush():
        if paragraph:
            value = " ".join(part.strip() for part in paragraph if part.strip())
            paragraph.clear()
            if value:
                return ("paragraph", _plain(value))
        return None

    in_code = False
    code: list[str] = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                yield ("code", "\n".join(code))
                code.clear()
            else:
                pending = flush()
                if pending:
                    yield pending
            in_code = not in_code
            continue
        if in_code:
            code.append(line)
            continue
        if not line.strip():
            pending = flush()
            if pending:
                yield pending
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        quote = re.match(r"^>\s?(.*)$", line)
        if heading or bullet or numbered or quote:
            pending = flush()
            if pending:
                yield pending
            if heading:
                yield (f"heading{len(heading.group(1))}", _plain(heading.group(2)))
            elif bullet:
                yield ("bullet", _plain(bullet.group(1)))
            elif numbered:
                yield ("number", _plain(numbered.group(1)))
            else:
                yield ("quote", _plain(quote.group(1)))
        else:
            paragraph.append(line)
    pending = flush()
    if pending:
        yield pending
    if code:
        yield ("code", "\n".join(code))


def render_docx(markdown: str) -> bytes:
    """Create a polished, editable Word document from academic Markdown."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.9)

    def set_font(style, size: float, color: str = "1F2937", bold=False):
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold

    normal = doc.styles["Normal"]
    set_font(normal, 11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.35
    for name, size, color in (
        ("Title", 24, "17365D"), ("Heading 1", 17, "17365D"),
        ("Heading 2", 14, "2E5C8A"), ("Heading 3", 12, "365F7D"),
    ):
        set_font(doc.styles[name], size, color, True)
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("科研助手 Agent  ·  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    first_heading = True
    for kind, value in _blocks(markdown):
        if kind.startswith("heading"):
            level = int(kind[-1])
            if first_heading and level == 1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(18)
                p.add_run(value)
                first_heading = False
            else:
                doc.add_paragraph(value, style=f"Heading {min(level, 3)}")
        elif kind == "bullet":
            doc.add_paragraph(value, style="List Bullet")
        elif kind == "number":
            doc.add_paragraph(value, style="List Number")
        elif kind == "quote":
            p = doc.add_paragraph(value)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.runs[0].italic = True
            p.runs[0].font.color.rgb = RGBColor(75, 85, 99)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(value)
            run.font.name, run.font.size = "Consolas", Pt(9)
        else:
            p = doc.add_paragraph(value)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.29)

    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def render_pdf(markdown: str) -> bytes:
    """Create an A4 PDF with Chinese-safe fonts and automatic page numbering."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font_name = "STSong-Light"
    try:
        pdfmetrics.getFont(font_name)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))

    base = getSampleStyleSheet()["BodyText"]
    body = ParagraphStyle("AcademicBody", parent=base, fontName=font_name, fontSize=10.5,
                          leading=17, spaceAfter=7, alignment=TA_JUSTIFY,
                          textColor=colors.HexColor("#1F2937"), firstLineIndent=21)
    title = ParagraphStyle("AcademicTitle", parent=body, fontSize=22, leading=30,
                           alignment=TA_CENTER, textColor=colors.HexColor("#17365D"),
                           spaceAfter=18, firstLineIndent=0)
    headings = {
        level: ParagraphStyle(f"H{level}", parent=body, fontSize=size, leading=size + 6,
                              spaceBefore=12 - level, spaceAfter=7 - level,
                              textColor=colors.HexColor(color), firstLineIndent=0)
        for level, size, color in (
            (1, 16, "#17365D"),
            (2, 13, "#2E5C8A"),
            (3, 11.5, "#365F7D"),
        )
    }
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=18, firstLineIndent=-10)
    quote = ParagraphStyle("Quote", parent=body, leftIndent=22, rightIndent=22,
                           textColor=colors.HexColor("#4B5563"), firstLineIndent=0)
    code = ParagraphStyle("Code", parent=body, fontSize=8.5, leading=12, leftIndent=14,
                          firstLineIndent=0, backColor=colors.HexColor("#F3F4F6"), borderPadding=6)

    story, first_heading, number = [], True, 0
    for kind, value in _blocks(markdown):
        safe = escape(value).replace("\n", "<br/>")
        if kind.startswith("heading"):
            level = int(kind[-1])
            if first_heading and level == 1:
                story.extend((Spacer(1, 14 * mm), Paragraph(safe, title)))
                first_heading = False
            else:
                story.append(Paragraph(safe, headings[min(level, 3)]))
        elif kind == "bullet":
            story.append(Paragraph(f"- {safe}", bullet))
        elif kind == "number":
            number += 1
            story.append(Paragraph(f"{number}. {safe}", bullet))
        elif kind == "quote":
            story.append(Paragraph(safe, quote))
        elif kind == "code":
            story.append(Paragraph(safe, code))
        else:
            number = 0
            story.append(Paragraph(safe, body))

    stream = BytesIO()
    pdf = SimpleDocTemplate(stream, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm,
                            topMargin=20 * mm, bottomMargin=18 * mm, title="科研写作导出",
                            author="科研助手 Agent")

    def page(canvas, document):
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawCentredString(A4[0] / 2, 9 * mm, f"科研助手 Agent  ·  {document.page}")
        canvas.restoreState()

    pdf.build(story or [Paragraph("（暂无正文）", body)], onFirstPage=page, onLaterPages=page)
    return stream.getvalue()
